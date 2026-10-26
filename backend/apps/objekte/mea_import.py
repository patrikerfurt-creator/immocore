"""MEA-Import je Liegenschaft — Parsing, Vorschau, Commit.

Liest einen MEA-Export (z. B. aus iX-Haus) und setzt die Miteigentumsanteile
auf den Verteilerschlüssel ``010`` (MEA, zeitlos / wirtschaftsjahr=0) der
bereits vorhandenen Einheiten eines Objekts.

Spaltenerwartung (nach Kopfzeilen-Erkennung, sonst feste Positionen):
    A  Objekt Nr.      → muss zum gewählten Objekt passen
    B  Fl.Nr.          → Zuordnung über ``Einheit.flaechennummer``
    F  Wert            → MEA-Wert (Dezimal, i. d. R. Zehntausendstel)

Es werden ausschließlich bestehende Einheiten aktualisiert — es werden keine
Einheiten angelegt. Nicht zuordenbare Zeilen erscheinen als Warnung in der
Vorschau und werden übersprungen.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from django.db import models, transaction

# Verteilerschlüssel-Code für MEA (siehe apps.konten.services.MUSTER_VS)
MEA_SCHLUESSEL = '010'
MEA_WIRTSCHAFTSJAHR = 0   # MEA ist zeitlos

# Feste Spaltenpositionen (0-basiert), falls die Kopfzeile nicht erkannt wird.
_FALLBACK_COL_OBJEKT = 0   # A
_FALLBACK_COL_FLNR   = 1   # B
_FALLBACK_COL_WERT   = 5   # F

_MAX_ZEILEN = 5000


# ---------------------------------------------------------------------------
# Rohdaten lesen (.xls via xlrd, .xlsx via openpyxl)
# ---------------------------------------------------------------------------
def _lese_tabelle(file_bytes: bytes, filename: str) -> list[list[Any]]:
    """Gibt alle Zeilen des ersten Blatts als Liste von Zell-Listen zurück."""
    name = (filename or '').lower()
    ist_xls = name.endswith('.xls')

    if ist_xls:
        try:
            import xlrd
        except ImportError:  # pragma: no cover - Abhängigkeit in requirements
            raise ValueError(
                'Für .xls-Dateien fehlt die Bibliothek "xlrd" auf dem Server.'
            )
        try:
            wb = xlrd.open_workbook(file_contents=file_bytes)
        except Exception as exc:
            raise ValueError(f'Datei konnte nicht gelesen werden: {exc}')
        sh = wb.sheet_by_index(0)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)]
                for r in range(sh.nrows)]

    # .xlsx / .xlsm
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        raise ValueError('Für .xlsx-Dateien fehlt die Bibliothek "openpyxl".')
    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f'Datei konnte nicht gelesen werden: {exc}')
    ws = wb[wb.sheetnames[0]]
    return [list(row) for row in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------------------
# Hilfsfunktionen zur Zell-Normalisierung
# ---------------------------------------------------------------------------
def _zelle_text(val: Any) -> str:
    """Zellwert als bereinigter Text. Ganzzahlige Floats ohne '.0'."""
    if val is None:
        return ''
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _normalisiere_objektnummer(val: Any) -> str:
    return _zelle_text(val)


def _parse_wert(val: Any) -> Decimal | None:
    """Wandelt einen Zellwert in einen MEA-Dezimalwert (4 Nachkommastellen)."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        dez = Decimal(str(val))
    else:
        s = str(val).strip()
        if not s:
            return None
        # Deutsches Format '1.234,56' → Punkt (Tausender) entfernen, Komma→Punkt.
        if ',' in s and '.' in s and s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '')
        s = s.replace(',', '.')
        try:
            dez = Decimal(s)
        except InvalidOperation:
            raise ValueError('ungültiger Zahlenwert')
    if dez < 0:
        raise ValueError('negativer Wert')
    return dez.quantize(Decimal('0.0001'))


# ---------------------------------------------------------------------------
# Kopfzeile erkennen und Spalten zuordnen
# ---------------------------------------------------------------------------
def _finde_spalten(tabelle: list[list[Any]]) -> tuple[int, dict[str, int]]:
    """Sucht die Kopfzeile und ordnet Objekt/Flächennummer/Wert zu.

    Gibt (kopfzeilen_index, {'objekt': i, 'flnr': j, 'wert': k}) zurück.
    Fällt auf feste Spalten A/B/F zurück, wenn keine Kopfzeile erkennbar ist.
    """
    for idx, zeile in enumerate(tabelle[:10]):
        texte = [_zelle_text(c).lower().replace(' ', '').rstrip('.') for c in zeile]
        col_obj = col_fl = col_wert = None
        for ci, t in enumerate(texte):
            if col_obj is None and t.startswith('objekt') and ('nr' in t or 'nummer' in t):
                col_obj = ci
            elif col_fl is None and (t in ('fl.nr', 'flnr', 'fl-nr', 'flächennr',
                                           'flaechennr', 'flächennummer', 'flaechennummer')
                                     or (t.startswith('fl') and 'nr' in t)):
                col_fl = ci
            elif col_wert is None and t in ('wert', 'mea', 'meawert', 'anteil', 'mea-wert'):
                col_wert = ci
        if col_fl is not None and col_wert is not None:
            return idx, {
                'objekt': col_obj if col_obj is not None else _FALLBACK_COL_OBJEKT,
                'flnr':   col_fl,
                'wert':   col_wert,
            }

    # Keine Kopfzeile erkannt → feste Positionen, Daten ab Zeile 2 annehmen.
    return 0, {
        'objekt': _FALLBACK_COL_OBJEKT,
        'flnr':   _FALLBACK_COL_FLNR,
        'wert':   _FALLBACK_COL_WERT,
    }


# ---------------------------------------------------------------------------
# Datei parsen → rohe Zeilen
# ---------------------------------------------------------------------------
def parse_mea_datei(file_bytes: bytes, filename: str) -> list[dict]:
    """Liest die Datei und gibt rohe Zeilen-Dicts zurück (ohne DB-Zugriff).

    Jede Zeile: {zeile, objektnummer, flaechennummer, wert, wert_fehler}
    """
    tabelle = _lese_tabelle(file_bytes, filename)
    if not tabelle:
        raise ValueError('Die Datei enthält keine Daten.')

    kopf_idx, cols = _finde_spalten(tabelle)
    max_col = max(cols.values())

    rows: list[dict] = []
    for r in range(kopf_idx + 1, len(tabelle)):
        zeile = tabelle[r]
        if len(zeile) <= max_col:
            zeile = list(zeile) + [None] * (max_col + 1 - len(zeile))

        flaechennummer = _zelle_text(zeile[cols['flnr']])
        objektnummer   = _normalisiere_objektnummer(zeile[cols['objekt']])
        wert_roh       = zeile[cols['wert']]

        # Vollständig leere Zeilen überspringen.
        if not flaechennummer and (wert_roh is None or _zelle_text(wert_roh) == '') \
                and not objektnummer:
            continue

        wert = None
        wert_fehler = None
        try:
            wert = _parse_wert(wert_roh)
        except ValueError as exc:
            wert_fehler = str(exc)

        rows.append({
            'zeile':          r + 1,   # 1-basiert für Anzeige
            'objektnummer':   objektnummer,
            'flaechennummer': flaechennummer,
            'wert':           wert,
            'wert_fehler':    wert_fehler,
        })

        if len(rows) > _MAX_ZEILEN:
            raise ValueError(f'Maximal {_MAX_ZEILEN} Zeilen pro Datei.')

    if not rows:
        raise ValueError('Keine Datenzeilen gefunden (Kopfzeile ohne Werte?).')
    return rows


# ---------------------------------------------------------------------------
# Vorschau erstellen (kein DB-Commit)
# ---------------------------------------------------------------------------
def mea_vorschau(objekt, file_bytes: bytes, filename: str) -> dict:
    """Ordnet die Datei-Zeilen den Einheiten des Objekts zu. Kein Schreibzugriff.

    Status je Zeile:
        ok             – zugeordnet, Wert gültig → wird importiert
        fremdes_objekt – Objektnummer passt nicht zum gewählten Objekt
        nicht_gefunden – Flächennummer nicht im Objekt vorhanden
        kein_wert      – kein/ungültiger Wert
        doppelt        – Flächennummer kommt in der Datei mehrfach vor
    """
    from .models import Einheit, Verteilerschluessel

    rows = parse_mea_datei(file_bytes, filename)

    objektnummer = (objekt.objektnummer or '').strip()

    # Flächennummer → Einheit (bestehende Einheiten des Objekts)
    einheit_map = {
        (e.flaechennummer or '').strip(): e
        for e in objekt.einheiten.all()
        if (e.flaechennummer or '').strip()
    }

    # aktuelle MEA-Werte je Einheit (für Alt→Neu-Anzeige)
    alt_werte: dict[str, Decimal] = {}
    vs = Verteilerschluessel.objects.filter(
        objekt=objekt, schluessel=MEA_SCHLUESSEL
    ).first()
    if vs is not None:
        for w in vs.werte.filter(wirtschaftsjahr=MEA_WIRTSCHAFTSJAHR).select_related('einheit'):
            alt_werte[str(w.einheit_id)] = w.wert

    gesehen: set[str] = set()
    zeilen = []
    objektnummern_datei: set[str] = set()

    for row in rows:
        flnr = row['flaechennummer']
        if row['objektnummer']:
            objektnummern_datei.add(row['objektnummer'])

        status = 'ok'
        meldung = ''
        einheit = einheit_map.get(flnr)

        if row['objektnummer'] and objektnummer and row['objektnummer'] != objektnummer:
            status = 'fremdes_objekt'
            meldung = (f'Objektnummer "{row["objektnummer"]}" gehört nicht zum '
                       f'gewählten Objekt "{objektnummer}".')
        elif not flnr:
            status = 'kein_wert'
            meldung = 'Keine Flächennummer in der Zeile.'
        elif flnr in gesehen:
            status = 'doppelt'
            meldung = f'Flächennummer "{flnr}" kommt in der Datei mehrfach vor.'
        elif einheit is None:
            status = 'nicht_gefunden'
            meldung = f'Keine Einheit mit Flächennummer "{flnr}" im Objekt.'
        elif row['wert_fehler'] is not None:
            status = 'kein_wert'
            meldung = f'Ungültiger Wert ({row["wert_fehler"]}).'
        elif row['wert'] is None:
            status = 'kein_wert'
            meldung = 'Kein Wert angegeben.'

        if flnr:
            gesehen.add(flnr)

        alt = alt_werte.get(str(einheit.id)) if einheit else None
        zeilen.append({
            'zeile':          row['zeile'],
            'status':         status,
            'meldung':        meldung,
            'flaechennummer': flnr,
            'objektnummer':   row['objektnummer'],
            'einheit_nr':     einheit.einheit_nr if einheit else '',
            'lage':           einheit.lage if einheit else '',
            'alt_wert':       str(alt) if alt is not None else '',
            'neu_wert':       str(row['wert']) if row['wert'] is not None else '',
            'daten': {
                'einheit_id': str(einheit.id) if einheit else None,
                'wert':       str(row['wert']) if row['wert'] is not None else None,
            },
        })

    ok_zeilen   = [z for z in zeilen if z['status'] == 'ok']
    summe       = sum((Decimal(z['neu_wert']) for z in ok_zeilen), Decimal('0'))
    zugeordnete_einheiten = {z['daten']['einheit_id'] for z in ok_zeilen}

    # Einheiten ohne passende Zeile in der Datei (reine Info)
    ohne_zeile = [
        {'flaechennummer': flnr, 'einheit_nr': e.einheit_nr}
        for flnr, e in einheit_map.items()
        if str(e.id) not in zugeordnete_einheiten
    ]

    summary = {
        'objektnummer':        objektnummer,
        'objektnummer_datei':  sorted(objektnummern_datei),
        'gesamt':              len(zeilen),
        'ok':                  len(ok_zeilen),
        'nicht_gefunden':      sum(1 for z in zeilen if z['status'] == 'nicht_gefunden'),
        'fremdes_objekt':      sum(1 for z in zeilen if z['status'] == 'fremdes_objekt'),
        'kein_wert':           sum(1 for z in zeilen if z['status'] == 'kein_wert'),
        'doppelt':             sum(1 for z in zeilen if z['status'] == 'doppelt'),
        'einheiten_gesamt':    len(einheit_map),
        'einheiten_ohne_zeile': ohne_zeile,
        'summe':               str(summe),
    }

    return {'zeilen': zeilen, 'summary': summary}


# ---------------------------------------------------------------------------
# Commit — MEA-Werte schreiben
# ---------------------------------------------------------------------------
@transaction.atomic
def mea_commit(objekt, rows: list[dict], user=None) -> dict:
    """Schreibt die MEA-Werte der übergebenen (geprüften) Zeilen.

    rows: Liste wie aus ``mea_vorschau`` → jede Zeile mit status und
    ``daten.einheit_id`` / ``daten.wert``. Es werden nur Zeilen mit
    status=='ok' geschrieben. Atomar.
    """
    from .models import Verteilerschluessel, VerteilerschluesselWert
    from apps.konten.services import verteilerschluessel_anlegen

    # Sicherstellen, dass der MEA-Schlüssel (010) und Wertzeilen existieren.
    verteilerschluessel_anlegen(objekt.id)

    vs = Verteilerschluessel.objects.filter(
        objekt=objekt, schluessel=MEA_SCHLUESSEL
    ).first()
    if vs is None:
        raise ValueError('MEA-Verteilerschlüssel (010) konnte nicht angelegt werden.')

    gueltige_einheit_ids = set(
        str(i) for i in objekt.einheiten.values_list('id', flat=True)
    )

    aktualisiert = 0
    uebersprungen = 0
    fehler: list[str] = []

    for row in rows:
        if row.get('status') != 'ok':
            uebersprungen += 1
            continue
        daten = row.get('daten') or {}
        einheit_id = daten.get('einheit_id')
        wert_roh   = daten.get('wert')
        if not einheit_id or einheit_id not in gueltige_einheit_ids or wert_roh is None:
            uebersprungen += 1
            continue
        try:
            wert = Decimal(str(wert_roh))
        except (InvalidOperation, ValueError):
            fehler.append(f'Zeile {row.get("zeile", "?")}: ungültiger Wert.')
            continue

        VerteilerschluesselWert.objects.update_or_create(
            schluessel=vs,
            einheit_id=einheit_id,
            wirtschaftsjahr=MEA_WIRTSCHAFTSJAHR,
            defaults={'wert': wert, 'beteiligt': True, 'quelle': 'stammdaten'},
        )
        aktualisiert += 1

    gesamt_summe = VerteilerschluesselWert.objects.filter(
        schluessel=vs, wirtschaftsjahr=MEA_WIRTSCHAFTSJAHR, wert__isnull=False,
    ).aggregate(s=models.Sum('wert'))['s'] or Decimal('0')

    return {
        'aktualisiert':  aktualisiert,
        'uebersprungen': uebersprungen,
        'fehler':        fehler,
        'summe':         str(gesamt_summe),
    }
